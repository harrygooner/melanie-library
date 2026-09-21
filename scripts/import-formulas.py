#!/usr/bin/env python3
"""Import Sapharchem formula workbooks/PDFs into a compact website dataset."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

import pdfplumber
from openpyxl import load_workbook


SOLUTION_DEFS = [
    {"id": "solution-moisture", "title": "Dưỡng ẩm", "description": "Công thức cấp ẩm, giữ ẩm và làm mềm da."},
    {"id": "solution-brightening", "title": "Làm sáng & nâng tông", "description": "Serum và kem hỗ trợ làm sáng, đều màu da."},
    {"id": "solution-antiaging", "title": "Chống lão hóa & săn chắc", "description": "Giải pháp đàn hồi, săn chắc và chăm sóc dấu hiệu tuổi tác."},
    {"id": "solution-soothing", "title": "Làm dịu & phục hồi", "description": "Công thức cho da nhạy cảm, phục hồi và giảm khó chịu."},
    {"id": "solution-exfoliation", "title": "Tẩy tế bào chết", "description": "Peeling, scrub và công thức làm sạch lớp sừng."},
    {"id": "solution-acne", "title": "Mụn & kháng khuẩn", "description": "Serum, gel và sản phẩm làm sạch hỗ trợ da mụn."},
    {"id": "solution-sun", "title": "Chống nắng", "description": "Kem chống nắng và giải pháp bảo vệ da khỏi UV."},
    {"id": "solution-cleansing", "title": "Làm sạch", "description": "Sữa rửa mặt, sữa tắm và tẩy trang."},
    {"id": "solution-hair", "title": "Tóc & da đầu", "description": "Dầu gội, dầu xả, tạo kiểu và chăm sóc da đầu."},
    {"id": "solution-lip", "title": "Môi & trang điểm", "description": "Son dưỡng, son màu, cushion và công thức trang điểm."},
    {"id": "solution-intimate", "title": "Chăm sóc phụ nữ", "description": "Dung dịch, bọt, xịt và ampoule chăm sóc vùng kín."},
    {"id": "solution-oral", "title": "Chăm sóc răng miệng", "description": "Công thức kem đánh răng và oral care."},
    {"id": "solution-base", "title": "Nền công thức đa dụng", "description": "Nền kem, lotion, gel và prototype đa ứng dụng."},
]


def clean(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).replace("\u00a0", " ").replace("\uf0b0", "°")
    return re.sub(r"\s+", " ", text).strip()


def normalize(value: Any) -> str:
    text = unicodedata.normalize("NFD", clean(value)).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"[^a-z0-9]+", " ", text.lower()).strip()


def name_key(value: Any) -> str:
    return normalize(value)


def canonical_supplier(value: str) -> str:
    return "Imderma Laboratories" if normalize(value) == "imdermalab" else value


def first_code(*values: str) -> str:
    for value in values:
        match = re.search(r"\b(?:CB|SPC)\s?\d{2}(?:-[A-Z0-9]+){1,7}\b", value, re.I)
        if match:
            return match.group(0).replace(" ", "").upper()
    return ""


def display_title_from_filename(path: Path) -> str:
    title = path.stem
    title = re.sub(r"^(?:CB|SPC)\d{2}(?:-[A-Z0-9]+){1,7}[_ -]*", "", title, flags=re.I)
    title = re.sub(r"(?:_?FINAL|_?UPDATED|\s*-?\s*Copy)(?:[_ .-]*\d{2,4})*$", "", title, flags=re.I)
    title = re.sub(r"[_]+", " ", title)
    title = re.sub(r"\s+", " ", title).strip(" -_")
    return title or path.stem


def valid_formula_file(path: Path) -> bool:
    value = normalize(path.name)
    blocked = ["stability", "cost", "formula bank", "tong hop", "meeting", "label", "~$"]
    return path.suffix.lower() in {".xlsx", ".pdf"} and not path.name.startswith("~$") and not any(term in value for term in blocked)


def header_mapping(row: list[Any]) -> dict[str, int]:
    result: dict[str, int] = {}
    for index, value in enumerate(row):
        key = normalize(value)
        if "phase" in key or key == "pha":
            result.setdefault("phase", index)
        elif "trade name" in key or "ten thuong mai" in key:
            result.setdefault("tradeName", index)
        elif "inci" in key:
            result.setdefault("inci", index)
        elif "function" in key or "chuc nang" in key:
            result.setdefault("function", index)
        elif "supplier" in key or "nha cung cap" in key:
            result.setdefault("supplier", index)
        elif "%" in clean(value) or "ham luong" in key or key == "w w":
            result.setdefault("percentage", index)
    return result


def ingredient_from_row(row: list[Any], columns: dict[str, int], previous_phase: str) -> tuple[dict[str, str] | None, str]:
    def at(key: str) -> str:
        index = columns.get(key, -1)
        return clean(row[index]) if 0 <= index < len(row) else ""

    phase = at("phase") or previous_phase
    trade = at("tradeName")
    inci = at("inci")
    function = at("function")
    supplier = at("supplier")
    percentage = at("percentage")
    try:
        number = float(percentage)
        percentage = f"{number:.4f}".rstrip("0").rstrip(".")
    except (TypeError, ValueError):
        pass
    marker = normalize(" ".join([phase, trade, inci]))
    if marker.startswith("tong cong") or marker.startswith("total") or marker.startswith("notice"):
        return None, phase
    if not trade and not inci:
        return None, phase
    if normalize(trade) in {"trade name", "ten thuong mai"}:
        return None, phase
    if normalize(inci) in {"inci name", "ten inci"}:
        return None, phase
    item = {
        "phase": phase,
        "tradeName": trade,
        "inci": inci,
        "function": function,
        "supplier": supplier,
        "percentage": percentage,
    }
    return item, phase


def xlsx_formulas(path: Path, year: int) -> list[dict[str, Any]]:
    formulas: list[dict[str, Any]] = []
    try:
        workbook = load_workbook(path, data_only=True, read_only=True)
    except Exception:
        return formulas
    for sheet in workbook.worksheets:
        if any(term in normalize(sheet.title) for term in ["cost", "price", "gia thanh"]):
            continue
        rows = list(sheet.iter_rows(values_only=True))
        header_index = -1
        columns: dict[str, int] = {}
        for index, row in enumerate(rows[:30]):
            mapping = header_mapping(list(row))
            if "tradeName" in mapping and "inci" in mapping and "percentage" in mapping:
                header_index, columns = index, mapping
                break
        if header_index < 0:
            continue
        heading_values = [clean(value) for row in rows[:header_index] for value in row if clean(value)]
        heading = " ".join(heading_values[:5])
        title = clean(heading_values[0]).split("Formula No.")[0].split("Code:")[0].strip() if heading_values else ""
        if not title or len(title) > 100:
            title = display_title_from_filename(path)
        code = first_code(heading, path.stem)
        ingredients: list[dict[str, str]] = []
        phase = ""
        empty_run = 0
        for row in rows[header_index + 1 :]:
            item, phase = ingredient_from_row(list(row), columns, phase)
            if item:
                ingredients.append(item)
                empty_run = 0
            else:
                empty_run += 1
                if ingredients and empty_run >= 8:
                    break
        if len(ingredients) < 2:
            continue
        if len(workbook.sheetnames) > 1 and sheet.title.lower() not in {"sheet1", "formula", "công thức"}:
            title = f"{title} — {sheet.title}"
        formulas.append({
            "title": title,
            "code": code,
            "year": year,
            "sourceFile": path.name,
            "sourceType": "XLSX",
            "ingredients": ingredients,
        })
    return formulas


def pdf_title_and_code(text: str, path: Path) -> tuple[str, str]:
    code = first_code(text[:700], path.stem)
    before = re.split(r"Formula\s*No\.?\s*:|Code\s*:", text[:1000], maxsplit=1, flags=re.I)[0]
    before = re.split(r"M[oô] ?t[aả]\s*s[aả]n\s*ph[aẩ]m\s*:|Description\s*:", before, maxsplit=1, flags=re.I)[0]
    lines = [clean(line) for line in before.splitlines() if clean(line)]
    title = " ".join(lines[:2]) if lines else display_title_from_filename(path)
    if not title or len(title) > 120 or normalize(title).startswith("phase"):
        title = display_title_from_filename(path)
    return title, code


def pdf_formulas(path: Path, year: int) -> list[dict[str, Any]]:
    try:
        with pdfplumber.open(path) as pdf:
            text = "\n".join(page.extract_text() or "" for page in pdf.pages)
            tables = [table for page in pdf.pages for table in page.extract_tables() if table]
    except Exception:
        return []
    title, code = pdf_title_and_code(text, path)
    best: tuple[dict[str, int], list[list[Any]]] | None = None
    for table in tables:
        for index, row in enumerate(table[:5]):
            mapping = header_mapping(list(row))
            if "tradeName" in mapping and "inci" in mapping and "percentage" in mapping:
                candidate = (mapping, table[index + 1 :])
                if best is None or len(candidate[1]) > len(best[1]):
                    best = candidate
                break
    if not best:
        return []
    columns, rows = best
    ingredients: list[dict[str, str]] = []
    phase = ""
    for row in rows:
        item, phase = ingredient_from_row(list(row), columns, phase)
        if item:
            ingredients.append(item)
    if len(ingredients) < 2:
        return []
    return [{
        "title": title,
        "code": code,
        "year": year,
        "sourceFile": path.name,
        "sourceType": "PDF",
        "ingredients": ingredients,
    }]


def pick_solution(title: str, ingredients: list[dict[str, str]]) -> str:
    title_text = normalize(title)
    title_rules = [
        ("solution-oral", ["toothpaste", "oral", "rang mieng"]),
        ("solution-intimate", ["feminine", "vung kin", "yzon", "y zone"]),
        ("solution-sun", ["sunscreen", "spf50", "chong nang"]),
        ("solution-hair", ["shampoo", "conditioner", "hair spray", "scalp", "dandruff", "toc", "da dau"]),
        ("solution-cleansing", ["make up remover", "facial wash", "body wash", "shower gel", "cleanser", "remover", "3in1", "wash gel", "sua rua", "sua tam"]),
        ("solution-lip", ["lipstick", "lip balm", "lip glow", "cushion", "make up", "moisturizing lipstick"]),
        ("solution-exfoliation", ["exfoliat", "peeling", "scrub", "tay te bao chet"]),
        ("solution-acne", ["acne", "acnes", "anti acnes", "tri mun"]),
        ("solution-brightening", ["whitening", "tone up", "bright", "lam sang", "nang tong"]),
        ("solution-antiaging", ["elasticity", "firming", "stretch mark", "stretchmark", "anti aging", "lao hoa", "san chac"]),
        ("solution-soothing", ["recovery", "soothing", "diaper", "rash", "tamanu", "hemorrhoid", "phuc hoi", "lam diu"]),
        ("solution-moisture", ["moistur", "hydrat", "sleeping mask", "eye mask", "butter cream", "duong am", "toner"]),
    ]
    for solution_id, terms in title_rules:
        if any(term in title_text for term in terms):
            return solution_id
    function_text = normalize(" ".join(item["function"] for item in ingredients))
    for solution_id, terms in [
        ("solution-brightening", ["lam sang", "nang tong", "trang da"]),
        ("solution-antiaging", ["lao hoa", "san chac", "dan hoi"]),
        ("solution-soothing", ["phuc hoi", "lam diu", "giam viem"]),
        ("solution-moisture", ["duong am", "cap am", "giu am"]),
    ]:
        if any(term in function_text for term in terms):
            return solution_id
    return "solution-base"


def catalog_indexes(catalog: dict[str, Any]) -> tuple[dict[str, list[dict[str, Any]]], dict[str, list[dict[str, Any]]]]:
    by_name: dict[str, list[dict[str, Any]]] = defaultdict(list)
    by_inci: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for item in catalog["ingredients"]:
        by_name[name_key(item["name"])].append(item)
        inci_key = normalize(item.get("inci", ""))
        if inci_key:
            by_inci[inci_key].append(item)
    return by_name, by_inci


def best_match(item: dict[str, str], by_name: dict[str, list[dict[str, Any]]], by_inci: dict[str, list[dict[str, Any]]]) -> dict[str, Any] | None:
    trade = name_key(item["tradeName"])
    supplier = normalize(item["supplier"])
    candidates = list(by_name.get(trade, [])) if trade and trade not in {"-", "water"} else []
    if not candidates and len(trade.split()) >= 2:
        containing: list[dict[str, Any]] = []
        for key, values in by_name.items():
            if key.startswith(f"{trade} "):
                containing.extend(values)
        if len(containing) <= 8:
            candidates = containing
    if not candidates:
        inci_key = normalize(item["inci"])
        inci_candidates = by_inci.get(inci_key, []) if len(inci_key) >= 5 else []
        if len(inci_candidates) == 1:
            candidates = list(inci_candidates)
        elif supplier:
            candidates = [candidate for candidate in inci_candidates if normalize(candidate["supplier"]) in supplier or supplier in normalize(candidate["supplier"])]
    if not candidates:
        return None
    if supplier:
        preferred = [candidate for candidate in candidates if normalize(candidate["supplier"]) in supplier or supplier in normalize(candidate["supplier"])]
        if preferred:
            candidates = preferred
    candidates.sort(key=lambda candidate: (abs(len(name_key(candidate["name"])) - len(trade)), candidate["name"]))
    return candidates[0]


def formula_similarity(a: dict[str, Any], b: dict[str, Any]) -> float:
    a_set = {normalize(item["tradeName"] or item["inci"]) for item in a["ingredients"] if normalize(item["tradeName"] or item["inci"])}
    b_set = {normalize(item["tradeName"] or item["inci"]) for item in b["ingredients"] if normalize(item["tradeName"] or item["inci"])}
    if not a_set or not b_set:
        return 0
    return len(a_set & b_set) / len(a_set | b_set)


def deduplicate(formulas: list[dict[str, Any]]) -> list[dict[str, Any]]:
    groups: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for formula in formulas:
        key = normalize(formula["code"] + " " + formula["title"])
        groups[key].append(formula)
    result: list[dict[str, Any]] = []
    for candidates in groups.values():
        chosen: list[dict[str, Any]] = []
        candidates.sort(key=lambda item: (item["year"], item["sourceType"] == "XLSX", len(item["ingredients"])), reverse=True)
        for candidate in candidates:
            duplicate_index = next((index for index, existing in enumerate(chosen) if formula_similarity(candidate, existing) >= 0.82), -1)
            if duplicate_index < 0:
                chosen.append(candidate)
            elif len(candidate["ingredients"]) > len(chosen[duplicate_index]["ingredients"]):
                chosen[duplicate_index] = candidate
        result.extend(chosen)
    return result


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("catalog", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()

    catalog = json.loads(args.catalog.read_text(encoding="utf-8"))
    by_name, by_inci = catalog_indexes(catalog)
    extracted: list[dict[str, Any]] = []
    for path in sorted(args.source.rglob("*")):
        if not path.is_file() or not valid_formula_file(path):
            continue
        year_match = re.search(r"20(?:23|24|25)", str(path))
        year = int(year_match.group(0)) if year_match else 0
        extracted.extend(xlsx_formulas(path, year) if path.suffix.lower() == ".xlsx" else pdf_formulas(path, year))

    formulas = deduplicate(extracted)
    for formula in formulas:
        formula["solutionId"] = pick_solution(formula["title"], formula["ingredients"])
        signature = normalize(formula["code"] + " " + formula["title"] + " " + formula["sourceFile"])
        formula["id"] = "formula-" + hashlib.sha1(signature.encode()).hexdigest()[:12]
        linked = 0
        for index, ingredient in enumerate(formula["ingredients"], start=1):
            ingredient["supplier"] = canonical_supplier(ingredient.get("supplier", ""))
            ingredient["order"] = index
            match = best_match(ingredient, by_name, by_inci)
            ingredient["productId"] = match["id"] if match else ""
            ingredient["catalogName"] = match["name"] if match else ""
            linked += int(bool(match))
        formula["linkedIngredientCount"] = linked
        formula["ingredientCount"] = len(formula["ingredients"])

    formulas.sort(key=lambda item: (next(index for index, value in enumerate(SOLUTION_DEFS) if value["id"] == item["solutionId"]), -item["year"], item["title"]))
    counts = Counter(item["solutionId"] for item in formulas)
    solutions = [{**solution, "formulaCount": counts[solution["id"]]} for solution in SOLUTION_DEFS if counts[solution["id"]]]
    payload = {
        "meta": {
            "formulaCount": len(formulas),
            "solutionCount": len(solutions),
            "ingredientLineCount": sum(item["ingredientCount"] for item in formulas),
            "linkedIngredientLineCount": sum(item["linkedIngredientCount"] for item in formulas),
            "sourceYears": sorted({item["year"] for item in formulas if item["year"]}),
        },
        "solutions": solutions,
        "formulas": formulas,
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(json.dumps(payload["meta"], ensure_ascii=False))
    print(json.dumps({solution["title"]: solution["formulaCount"] for solution in solutions}, ensure_ascii=False))


if __name__ == "__main__":
    main()
