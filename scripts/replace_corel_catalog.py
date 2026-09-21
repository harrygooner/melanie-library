#!/usr/bin/env python3
"""Replace the Corel catalog slice with the grouped Corel workbook."""

from __future__ import annotations

import hashlib
import json
import re
import sys
import unicodedata
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


SUPPLIER = "Corel"

CATEGORY_GROUPS = {
    "ACRYPOL® – RHEOLOGY MODIFIERS": "Hệ nền – Chất điều chỉnh lưu biến / Tạo đặc",
    "ACRYPOL® – CONDITIONING RANGE": "Tóc & da đầu – Polymer conditioning / Chống tĩnh điện",
    "ACRYPOL® CP5 – DETERGENCY / ANTI-REDEPOSITION": "Home Care – Polymer hỗ trợ tẩy rửa / Chống tái lắng đọng",
    "ACRYM® – MULTIFUNCTIONAL 3-IN-1 LIQUID POLYMERS": "Hệ nền – Polymer lỏng đa chức năng",
    "ACRYSET® – HAIR FIXATIVE POLYMERS": "Tóc & da đầu – Polymer tạo màng / Giữ nếp",
    "ACRYFILM® – FLEXIBLE FILM FORMERS": "Hệ nền – Polymer tạo màng linh hoạt",
    "ACRYSOL® – SOLUBILIZERS & EMULSIFYING AGENTS": "Hệ nền – Chất hòa tan / Hỗ trợ nhũ hóa",
}


def clean(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value).replace("\u00a0", " ")).strip()


def normalize(value: Any) -> str:
    text = unicodedata.normalize("NFKD", clean(value))
    text = "".join(char for char in text if not unicodedata.combining(char))
    return re.sub(r"[^a-z0-9]+", " ", text.casefold().replace("đ", "d")).strip()


def name_key(value: Any) -> str:
    text = normalize(value)
    text = re.sub(r"\bkhong chua benzene\b", "", text)
    text = re.sub(r"\bdong benzene free.*$", "", text)
    return re.sub(r"\s+", " ", text).strip()


def product_id(name: str) -> str:
    digest = hashlib.sha1(f"{SUPPLIER}|{name}".encode("utf-8")).hexdigest()[:12]
    return f"spc-{digest}"


def parse_corel(workbook_path: Path, old_items: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], dict[str, list[str]]]:
    workbook = load_workbook(workbook_path, read_only=False, data_only=True)
    sheet = workbook["Corel Portfolio"]

    old_by_key: dict[str, dict[str, Any]] = {}
    for item in old_items:
        old_by_key[name_key(item["name"])] = item

    # These former combined catalog rows are split into individual grades in the new workbook.
    split_ids = {
        "acrypol 934": "spc-38a0106fdcc2",
        "acrypol 940": "spc-e4acd84282ff",
    }
    aliases = {
        "acrypol et 28": "acrypol et 28",
        "acrypol et2020": "acrypol et2020",
        "acrypol elt 10": "acrypol elt 10",
        "acrypol elt 20": "acrypol elt 20",
        "acrypol elt 21": "acrypol elt 21",
        "acryset superhold ii": "acryset superhold ii",
        "acryset p 30": "acryset p 30",
        "acrysol k 140": "acrysol k 140",
    }

    current_group = ""
    new_items: list[dict[str, Any]] = []
    old_to_new: dict[str, list[str]] = {}
    for row in sheet.iter_rows(values_only=True):
        first = clean(row[0] if row else "")
        if first in CATEGORY_GROUPS:
            current_group = CATEGORY_GROUPS[first]
            continue
        if len(row) < 2 or not isinstance(row[0], (int, float)) or not clean(row[1]):
            continue

        name = clean(row[1])
        form = clean(row[2])
        inci = clean(row[3])
        description = clean(row[4])
        certificates = clean(row[5])
        key = name_key(name)
        key = aliases.get(key, key)
        existing = old_by_key.get(key)
        pid = existing["id"] if existing else split_ids.get(key, product_id(name))
        if existing:
            old_to_new.setdefault(existing["id"], []).append(pid)
        if key in split_ids:
            old_to_new.setdefault(split_ids[key], []).append(pid)

        new_items.append(
            {
                "id": pid,
                "group": current_group,
                "supplier": SUPPLIER,
                "name": name,
                "form": form,
                "inci": inci,
                "description": description,
                "certificates": certificates,
                "certificateTags": [],
                "search": normalize(" ".join([current_group, SUPPLIER, name, form, inci, description, certificates])),
            }
        )

    # Explicitly split the two legacy combined rows across both new grades.
    by_key = {name_key(item["name"]): item["id"] for item in new_items}
    old_to_new["spc-38a0106fdcc2"] = [by_key["acrypol 934"], by_key["acrypol 974"]]
    old_to_new["spc-e4acd84282ff"] = [by_key["acrypol 940"], by_key["acrypol 980"]]
    return new_items, old_to_new


def recalculate_catalog(catalog: dict[str, Any], workbook_path: Path) -> None:
    ingredients = catalog["ingredients"]
    groups = Counter(item["group"] for item in ingredients)
    suppliers = Counter(item["supplier"] for item in ingredients)
    certificates = Counter(tag for item in ingredients for tag in item["certificateTags"])
    for solution in catalog["packedSolutions"]:
        solution["productIds"] = list(dict.fromkeys(solution["productIds"]))
        solution["count"] = len(solution["productIds"])

    workbook = load_workbook(workbook_path, read_only=False, data_only=True)
    modified = workbook.properties.modified.isoformat() if workbook.properties.modified else ""
    catalog["meta"].update(
        {
            "sourceModified": modified,
            "ingredientCount": len(ingredients),
            "groupCount": len(groups),
            "supplierCount": len(suppliers),
            "packedSolutionCount": len(catalog["packedSolutions"]),
            "corelSourceFile": workbook_path.name,
            "corelSourceModified": modified,
        }
    )
    catalog["groups"] = [
        {"name": name, "count": count}
        for name, count in sorted(groups.items(), key=lambda item: (-item[1], item[0]))
    ]
    catalog["suppliers"] = [
        {"name": name, "count": count}
        for name, count in sorted(suppliers.items(), key=lambda item: (-item[1], item[0]))
    ]
    catalog["certificates"] = [
        {"name": name, "count": count}
        for name, count in sorted(certificates.items(), key=lambda item: (-item[1], item[0]))
    ]


def main() -> None:
    if len(sys.argv) != 4:
        raise SystemExit("Usage: replace_corel_catalog.py COREL.xlsx CATALOG.json FORMULAS.json")
    workbook_path = Path(sys.argv[1]).resolve()
    catalog_path = Path(sys.argv[2]).resolve()
    formulas_path = Path(sys.argv[3]).resolve()

    catalog = json.loads(catalog_path.read_text(encoding="utf-8"))
    formulas = json.loads(formulas_path.read_text(encoding="utf-8"))
    old_items = [item for item in catalog["ingredients"] if normalize(item["supplier"]) == "corel"]
    new_items, old_to_new = parse_corel(workbook_path, old_items)

    catalog["ingredients"] = [
        item for item in catalog["ingredients"] if normalize(item["supplier"]) != "corel"
    ] + new_items
    new_ids = {item["id"] for item in new_items}

    for solution in catalog["packedSolutions"]:
        remapped: list[str] = []
        for pid in solution["productIds"]:
            if pid in old_to_new:
                remapped.extend(old_to_new[pid])
            elif pid not in {item["id"] for item in old_items}:
                remapped.append(pid)
        solution["productIds"] = [pid for pid in dict.fromkeys(remapped) if pid in {item["id"] for item in catalog["ingredients"]}]

    catalog_names = {item["id"]: item["name"] for item in new_items}
    for formula in formulas["formulas"]:
        for ingredient in formula["ingredients"]:
            pid = ingredient.get("productId", "")
            if pid in catalog_names:
                ingredient["catalogName"] = catalog_names[pid]
            elif pid in {item["id"] for item in old_items} and pid not in new_ids:
                ingredient["productId"] = ""
                ingredient["catalogName"] = ""
        formula["linkedIngredientCount"] = sum(1 for item in formula["ingredients"] if item.get("productId"))
    formulas["meta"]["linkedIngredientLineCount"] = sum(
        formula["linkedIngredientCount"] for formula in formulas["formulas"]
    )

    recalculate_catalog(catalog, workbook_path)
    catalog_path.write_text(json.dumps(catalog, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    formulas_path.write_text(json.dumps(formulas, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")

    print(
        json.dumps(
            {
                "removedCorel": len(old_items),
                "addedCorel": len(new_items),
                "ingredientCount": catalog["meta"]["ingredientCount"],
                "groupCount": catalog["meta"]["groupCount"],
                "supplierCount": catalog["meta"]["supplierCount"],
                "linkedFormulaLines": formulas["meta"]["linkedIngredientLineCount"],
            },
            ensure_ascii=False,
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
