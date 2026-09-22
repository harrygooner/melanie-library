from __future__ import annotations

import hashlib
import json
import re
import sys
import unicodedata
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


CERTIFICATE_PATTERNS = {
    "COSMOS": ("COSMOS",),
    "ECOCERT": ("ECOCERT",),
    "HALAL": ("HALAL", "HALAH"),
    "VEGAN": ("VEGAN", "VEGEAN"),
    "KOSHER": ("KOSHER",),
    "RSPO": ("RSPO",),
    "NATRUE": ("NATRUE",),
    "REACH": ("REACH",),
    "NATURAL": ("NATURAL", "TỰ NHIÊN"),
    "ORGANIC": ("ORGANIC",),
    "ISO": ("ISO",),
    "PATENTED": ("PATENT",),
}


def clean(value: object) -> str:
    if value is None:
        return ""
    return re.sub(r"[ \t]+", " ", str(value).replace("\r\n", "\n")).strip()


def normalize(value: str) -> str:
    value = unicodedata.normalize("NFKD", value)
    value = "".join(char for char in value if not unicodedata.combining(char))
    value = value.casefold().replace("đ", "d")
    return re.sub(r"[^a-z0-9]+", " ", value).strip()


def product_id(supplier: str, name: str) -> str:
    stable_supplier = "Imdermalab" if normalize(supplier) == "imderma laboratories" else supplier
    digest = hashlib.sha1(f"{stable_supplier}|{name}".encode("utf-8")).hexdigest()[:12]
    return f"spc-{digest}"


def canonical_supplier(value: str) -> str:
    return "Imderma Laboratories" if normalize(value) == "imdermalab" else value


def certificate_tags(raw: str) -> list[str]:
    upper = raw.upper()
    return [
        label
        for label, patterns in CERTIFICATE_PATTERNS.items()
        if any(pattern in upper for pattern in patterns)
    ]


def parse_catalog(workbook_path: Path) -> dict:
    wb = load_workbook(workbook_path, read_only=True, data_only=True)
    ws = wb["TỔNG HỢP"]

    ingredients: list[dict] = []
    index: dict[tuple[str, str], str] = {}
    for row in ws.iter_rows(min_row=4, values_only=True):
        group, supplier, name, form, inci, description, certificates = [
            clean(value) for value in row[:7]
        ]
        supplier = canonical_supplier(supplier)
        if not supplier or not name:
            continue
        key = (normalize(supplier), normalize(name))
        if key in index:
            continue
        pid = product_id(supplier, name)
        index[key] = pid
        tags = certificate_tags(certificates)
        ingredients.append(
            {
                "id": pid,
                "group": group,
                "supplier": supplier,
                "name": name,
                "form": form,
                "inci": inci,
                "description": description,
                "certificates": certificates,
                "certificateTags": tags,
                "search": normalize(
                    " ".join(
                        [group, supplier, name, form, inci, description, certificates]
                    )
                ),
            }
        )

    packed_ws = wb["PACKED SOLUTION"]
    packed: list[dict] = []
    current: dict | None = None
    missing = Counter()
    for row in packed_ws.iter_rows(values_only=True):
        values = [clean(value) for value in row[:7]]
        first = values[0]
        if first.upper().startswith("GIẢI PHÁP"):
            title = re.sub(r"\s*\(\d+\s+nguyên liệu\)\s*$", "", first, flags=re.I)
            title = re.sub(r"^GIẢI PHÁP\s+", "", title, flags=re.I).strip()
            current = {
                "id": f"packed-{len(packed) + 1}",
                "title": title,
                "productIds": [],
            }
            packed.append(current)
            continue
        if current is None or first == "NHÀ CUNG CẤP":
            continue
        supplier, name = values[0], values[1]
        if not supplier or not name:
            continue
        pid = index.get((normalize(supplier), normalize(name)))
        if pid:
            current["productIds"].append(pid)
        else:
            missing[(supplier, name)] += 1

    groups = Counter(item["group"] for item in ingredients)
    suppliers = Counter(item["supplier"] for item in ingredients)
    certificates = Counter(
        tag for item in ingredients for tag in item["certificateTags"]
    )
    for solution in packed:
        solution["productIds"] = list(dict.fromkeys(solution["productIds"]))
        solution["count"] = len(solution["productIds"])

    modified = wb.properties.modified.isoformat() if wb.properties.modified else ""
    return {
        "meta": {
            "sourceFile": workbook_path.name,
            "sourceModified": modified,
            "ingredientCount": len(ingredients),
            "groupCount": len(groups),
            "supplierCount": len(suppliers),
            "packedSolutionCount": len(packed),
            "unmatchedPackedRows": sum(missing.values()),
        },
        "groups": [
            {"name": name, "count": count}
            for name, count in sorted(groups.items(), key=lambda item: (-item[1], item[0]))
        ],
        "suppliers": [
            {"name": name, "count": count}
            for name, count in sorted(suppliers.items(), key=lambda item: (-item[1], item[0]))
        ],
        "certificates": [
            {"name": name, "count": count}
            for name, count in sorted(certificates.items(), key=lambda item: (-item[1], item[0]))
        ],
        "packedSolutions": packed,
        "ingredients": ingredients,
    }


def main() -> None:
    if len(sys.argv) != 3:
        raise SystemExit("Usage: extract_catalog.py INPUT.xlsx OUTPUT.json")
    source = Path(sys.argv[1]).resolve()
    output = Path(sys.argv[2]).resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    payload = parse_catalog(source)
    output.write_text(
        json.dumps(payload, ensure_ascii=False, separators=(",", ":")),
        encoding="utf-8",
    )
    print(json.dumps(payload["meta"], ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
